from pathlib import Path
import sys

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from exporter import export_result
from audit_logic import run_audit


def test_exports_compact_workbook(tmp_path):
    result = run_audit([
        {'order_no': 'A', 'seller': 'S', 'buyer_id': 'B', 'platform_sku': 'P', 'warehouse_sku': 'W', 'qty': 1},
    ], {'S': [
        {'seller': 'S', 'platform_sku': 'P', 'warehouse_sku': 'W', 'circle_stock': 1, 'distribution_stock': 0, 'cost_price': 1, 'discount_price': 1},
    ]})
    output = export_result(result, tmp_path / 'result.xlsx')
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == ['审单结果', '审单汇总']
    assert workbook['审单结果'].max_row == 2
    assert workbook['审单结果']['M2'].value == '可放行'
