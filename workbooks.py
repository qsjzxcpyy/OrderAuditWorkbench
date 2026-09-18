from __future__ import annotations

from pathlib import Path
from typing import Any
import re

import openpyxl


class WorkbookDataError(ValueError):
    pass


def _key(value: Any) -> str:
    text = str(value or '').strip().lower()
    return ''.join(ch for ch in text if not ch.isspace() and ch not in '_-（）()：:')


def _text(value: Any) -> str:
    return str(value or '').strip()


def _number(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    return _text(value)

ERP_ALIASES = {
    'order_no': {'refrencenoplatform', 'refrenceno', 'orderno', '订单号', '亚马逊订单号'},
    'seller': {'platformsellerid', '来源店铺', '店铺', 'seller'},
    'buyer_id': {'buyerid', 'platformbuyerid', '购买者id', '购买者ID'},
    'date_latest_ship': {'datelatestship'},
    'order_status': {'orderstatus'},
    'custom_order_type': {'customordertype'},
    'mark_content': {'markcontent'},
    'system_note': {'systemnote'},
    'operator_note': {'operatornote'},
    'abnormal_reason': {'abnormalreason'},
    'order_desc': {'orderdesc'},
    'amazon_order_remarks': {'amazonorderremarks'},
    'amazon_customized_order_notes': {'amazoncustomizedordernotes'},
    'warning_letter': {'warningletter'},
    'remark1': {'remark1'},
    'remark2': {'remark2'},
}

# Some exports label the child/platform SKU column as "?SKU"; the supplied
# THRYVIX workbook has the same column serialized as the middle-dot marker "?".
INVENTORY_ALIASES = {
    'platform_sku': {'\u4e9a\u9a6c\u900a\u5e73\u53f0sku', '\u4e9a\u9a6c\u900a\u5e73\u53f0SKU', '\u4ea7\u54c1SKU', '\u4ea7\u54c1sku', '\u5b50SKU', '\u5b50sku', '\u00b7', 'amazonsku', 'platformsku', 'productsku', '\u5e73\u53f0sku'},
    'warehouse_sku': {'云仓发货sku', '云仓发货SKU', 'warehousesku', '云仓sku'},
    'circle_stock': {'圈买库存', '圈货库存', '圈买库存数量'},
    'distribution_stock': {'分销库存', '云仓库存', '分销库存数量'},
    'cost_price': {'成本价', '成本价格', 'costprice'},
    'discount_price': {'当前分销优惠价价', '当前分销优惠价', '分销优惠价', 'discountprice'},
    'title': {'产品名称', '商品名称', 'productname'},
}


def _find_header(ws, aliases: dict[str, set[str]], scan_rows: int = 30) -> tuple[int, dict[str, int]]:
    normalized_aliases = {field: {_key(alias) for alias in values} for field, values in aliases.items()}
    best: tuple[int, dict[str, int]] | None = None
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=scan_rows, values_only=True), 1):
        columns: dict[str, int] = {}
        for index, value in enumerate(row):
            value_key = _key(value)
            for field, field_aliases in normalized_aliases.items():
                if value_key in field_aliases and field not in columns:
                    columns[field] = index
        if best is None or len(columns) > len(best[1]):
            best = (row_number, columns)
    if best is None:
        raise WorkbookDataError('无法找到工作表表头。')
    return best


def _sheet(workbook, sheet_name: str | None, aliases: dict[str, set[str]]):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise WorkbookDataError(f'找不到工作表：{sheet_name}')
        return workbook[sheet_name]
    best = None
    for name in workbook.sheetnames:
        ws = workbook[name]
        try:
            header = _find_header(ws, aliases)
        except WorkbookDataError:
            continue
        if best is None or len(header[1]) > len(best[1][1]):
            best = (name, header)
    if best is None:
        raise WorkbookDataError('无法找到包含所需字段的工作表。')
    return workbook[best[0]]


def _open(path: str | Path):
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookDataError(f'无法读取 Excel 文件：{Path(path).name}') from exc


def _parse_erp_workbook(workbook, sheet_name: str | None) -> list[dict[str, Any]]:
    ws = _sheet(workbook, sheet_name, ERP_ALIASES)
    header_row, header = _find_header(ws, ERP_ALIASES)
    missing = [field for field in ('order_no', 'seller', 'buyer_id') if field not in header]
    if missing:
        raise WorkbookDataError('ERP 文件缺少字段：' + '、'.join(missing))
    all_headers = [str(value or '') for value in next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
    # Preserve ERP descriptive fields so downstream rules can identify pre-sale
    # orders. These columns are intentionally explicit: product titles and address
    # fields must not accidentally turn a normal order into a pre-sale order.
    presale_fields = (
        'order_status', 'custom_order_type', 'mark_content', 'system_note',
        'operator_note', 'abnormal_reason', 'order_desc', 'amazon_order_remarks',
        'amazon_customized_order_notes', 'warning_letter', 'remark1', 'remark2',
    )
    presale_value_columns = [
        (field, header[field]) for field in presale_fields
        if field in header
    ]
    sku_columns: list[tuple[int, int, int]] = []
    for index, name in enumerate(all_headers):
        match = re.fullmatch(r'warehousesku(\d+)', _key(name))
        if match:
            item_number = int(match.group(1))
            sku_index = next((i for i, candidate in enumerate(all_headers) if _key(candidate) == f'sku{item_number}'), None)
            qty_index = next((i for i, candidate in enumerate(all_headers) if _key(candidate) == f'qty{item_number}'), None)
            if sku_index is not None and qty_index is not None:
                sku_columns.append((sku_index, index, qty_index))
    if not sku_columns:
        raise WorkbookDataError('ERP 文件缺少商品字段：sku1、warehouse_sku1、qty1')

    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        order_no = _text(values[header['order_no']] if header['order_no'] < len(values) else '')
        if not order_no:
            continue
        seller = _text(values[header['seller']] if header['seller'] < len(values) else '')
        buyer_id = _text(values[header['buyer_id']] if header['buyer_id'] < len(values) else '')
        latest_ship = _text(values[header['date_latest_ship']] if 'date_latest_ship' in header and header['date_latest_ship'] < len(values) else '')
        presale_values = [
            _text(values[index]) for _, index in presale_value_columns
            if index < len(values) and values[index] not in (None, '')
        ]
        descriptive_values = {
            field: _text(values[index])
            for field, index in presale_value_columns
            if index < len(values) and values[index] not in (None, '')
        }
        for sku_index, warehouse_index, qty_index in sku_columns:
            platform_sku = _text(values[sku_index] if sku_index < len(values) else '')
            warehouse_sku = _text(values[warehouse_index] if warehouse_index < len(values) else '')
            if not platform_sku and not warehouse_sku:
                continue
            rows.append({
                'order_no': order_no,
                'seller': seller,
                'buyer_id': buyer_id,
                'date_latest_ship': latest_ship,
                'presale_values': presale_values,
                **descriptive_values,
                'platform_sku': platform_sku,
                'warehouse_sku': warehouse_sku,
                'qty': _number(values[qty_index] if qty_index < len(values) else None),
            })
    return rows


def parse_erp_workbook(path: str | Path, sheet_name: str | None = None) -> list[dict[str, Any]]:
    workbook = _open(path)
    try:
        return _parse_erp_workbook(workbook, sheet_name)
    finally:
        workbook.close()


def _parse_inventory_workbook(workbook, path: str | Path, seller: str, sheet_name: str | None) -> list[dict[str, Any]]:
    ws = _sheet(workbook, sheet_name, INVENTORY_ALIASES)
    header_row, header = _find_header(ws, INVENTORY_ALIASES)
    required = ('platform_sku', 'warehouse_sku', 'circle_stock', 'distribution_stock', 'cost_price', 'discount_price')
    missing = [field for field in required if field not in header]
    if missing:
        raise WorkbookDataError(f'{Path(path).name} 缺少字段：' + '、'.join(missing))
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
        platform_sku = _text(values[header['platform_sku']] if header['platform_sku'] < len(values) else '')
        warehouse_sku = _text(values[header['warehouse_sku']] if header['warehouse_sku'] < len(values) else '')
        if not platform_sku and not warehouse_sku:
            continue
        rows.append({
            'seller': seller,
            'platform_sku': platform_sku,
            'warehouse_sku': warehouse_sku,
            'circle_stock': _number(values[header['circle_stock']] if header['circle_stock'] < len(values) else None),
            'distribution_stock': _number(values[header['distribution_stock']] if header['distribution_stock'] < len(values) else None),
            'cost_price': _number(values[header['cost_price']] if header['cost_price'] < len(values) else None),
            'discount_price': _number(values[header['discount_price']] if header['discount_price'] < len(values) else None),
            'title': _text(values[header['title']] if 'title' in header and header['title'] < len(values) else ''),
        })
    return rows


def parse_inventory_workbook(path: str | Path, seller: str, sheet_name: str | None = None) -> list[dict[str, Any]]:
    workbook = _open(path)
    try:
        return _parse_inventory_workbook(workbook, path, seller, sheet_name)
    finally:
        workbook.close()
