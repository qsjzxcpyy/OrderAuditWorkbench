from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


RESULT_HEADERS = [
    ('order_no', '订单号'),
    ('seller', '来源店铺'),
    ('buyer_id', 'buyer_id'),
    ('platform_sku', '亚马逊平台 SKU'),
    ('warehouse_sku', '云仓发货 SKU'),
    ('qty', '商品数量'),
    ('demand', '同云仓 SKU 当天总需求'),
    ('circle_stock', '圈买库存'),
    ('distribution_stock', '分销库存'),
    ('cost_price', '成本价'),
    ('discount_price', '当前分销优惠价'),
    ('price_difference', '价格差值'),
    ('status_text', '审单状态'),
    ('reasons_text', '人工处理原因'),
    ('duplicate_orders_text', '相关重复订单号'),
    ('audited_at', '审单时间'),
]


def _value(line: dict[str, Any], key: str, sku_groups: dict[str, Any]) -> Any:
    if key == 'demand':
        return sku_groups.get(line.get('warehouse_sku'), {}).get('demand')
    if key == 'status_text':
        return '人工处理' if line.get('status') == 'manual' else '可放行'
    if key == 'reasons_text':
        return '；'.join(line.get('reasons') or [])
    if key == 'duplicate_orders_text':
        return ' '.join(line.get('duplicate_orders') or [])
    if key == 'audited_at':
        return datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    if key in line:
        return line[key]
    return (line.get('evidence') or {}).get(key)


def export_result(result: dict[str, Any], output_path: str | Path) -> Path:
    workbook = Workbook()
    detail = workbook.active
    detail.title = '审单结果'
    summary = workbook.create_sheet('审单汇总')

    header_fill = PatternFill('solid', fgColor='13233A')
    header_font = Font(name='Arial', color='FFFFFF', bold=True)
    body_font = Font(name='Arial', color='1E2E42')

    detail.append([label for _, label in RESULT_HEADERS])
    for cell in detail[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for line in result.get('lines', []):
        row = [_value(line, key, result.get('sku_groups', {})) for key, _ in RESULT_HEADERS]
        detail.append(row)
    for row in detail.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    detail.freeze_panes = 'A2'
    detail.auto_filter.ref = detail.dimensions
    widths = [22, 18, 18, 24, 22, 12, 18, 12, 12, 12, 16, 12, 12, 44, 24, 20]
    for index, width in enumerate(widths, 1):
        detail.column_dimensions[get_column_letter(index)].width = width

    summary.append(['指标', '数值'])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary_rows = [
        ('总订单数', result.get('summary', {}).get('total_orders', 0)),
        ('可放行订单数', result.get('summary', {}).get('pass_orders', 0)),
        ('人工处理订单数', result.get('summary', {}).get('manual_orders', 0)),
        ('商品明细数', result.get('summary', {}).get('total_lines', 0)),
        ('审单时间', datetime.now().strftime('%Y-%m-%d %H:%M:%S')),
    ]
    for row in summary_rows:
        summary.append(row)
    for row in summary.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    summary.column_dimensions['A'].width = 24
    summary.column_dimensions['B'].width = 24
    summary.freeze_panes = 'A2'

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
